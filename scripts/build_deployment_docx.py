#!/usr/bin/env python3
"""Build the current Windows/Jetson deployment manual from its Markdown source."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_east_asia_font(style, name: str) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def shade(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def set_cell_margins(paragraph, before: int = 100, after: int = 100) -> None:
    properties = paragraph._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    properties.append(spacing)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    set_east_asia_font(normal, "Microsoft JhengHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    title = document.styles["Title"]
    set_east_asia_font(title, "Microsoft JhengHei")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(16)

    for name, size, before, after in (
        ("Heading 1", 17, 16, 7),
        ("Heading 2", 13, 12, 5),
        ("Heading 3", 11, 9, 4),
    ):
        style = document.styles[name]
        set_east_asia_font(style, "Microsoft JhengHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    set_east_asia_font(code, "Consolas")
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor(20, 20, 20)
    code.paragraph_format.left_indent = Cm(0.35)
    code.paragraph_format.right_indent = Cm(0.35)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)


def add_markdown(document: Document, source: str) -> None:
    in_code = False
    code_lines: list[str] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            document.add_paragraph(" ".join(line.strip() for line in paragraph_lines))
            paragraph_lines.clear()

    def flush_code() -> None:
        if code_lines:
            paragraph = document.add_paragraph("\n".join(code_lines), style="Code Block")
            shade(paragraph, "F2F2F2")
            set_cell_margins(paragraph)
            code_lines.clear()

    for raw_line in source.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            flush_paragraph()
            if in_code:
                flush_code()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line:
            flush_paragraph()
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                document.add_paragraph(text, style="Title")
                document.add_paragraph(
                    "本手冊說明兩個正式平台的安裝、模擬、實機安全與驗收流程。"
                    "正式系統統一使用 ROS 2 Humble，並以不可變映像 digest 部署。"
                )
            else:
                document.add_heading(text, level=level - 1)
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            document.add_paragraph(numbered.group(1), style="List Number")
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            document.add_paragraph(bullet.group(1), style="List Bullet")
            continue

        paragraph_lines.append(line)

    flush_paragraph()
    flush_code()


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    source = root / "docs" / "DEPLOYMENT_WINDOWS_JETSON.md"
    output = root / "docs" / "PhantomX拼圖系統_Windows_Jetson部署手冊.docx"

    document = Document()
    configure(document)
    add_markdown(document, source.read_text(encoding="utf-8"))
    document.core_properties.title = "Windows 與 Jetson 完整部署手冊"
    document.core_properties.subject = "PhantomX 拼圖系統 ROS 2 Humble 部署"
    document.core_properties.author = "ping152"
    document.core_properties.keywords = "ROS 2 Humble, PhantomX, Jetson, Windows, Docker"
    document.save(output)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
